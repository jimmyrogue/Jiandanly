from __future__ import annotations

import os
import shutil
import tempfile
from collections.abc import Callable
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from langchain_core.runnables.config import ensure_config
from openpyxl import load_workbook
from pptx import Presentation as PptxPresentation

from ..runtime import current_runtime_tool_execution

MARKDOWN_CHAR_CAP = 60_000
EXTENSION_KIND = {
    ".docx": "word",
    ".xlsx": "excel",
    ".pptx": "powerpoint",
}
EDITED_INFIX = "edited"


def workspace_root_from_config() -> str | None:
    try:
        config = ensure_config()
    except Exception:
        return None
    configurable = config.get("configurable") if isinstance(config, dict) else None
    if not isinstance(configurable, dict):
        return None
    if "workspace_root" not in configurable:
        return None
    return str(configurable.get("workspace_root") or "").strip()


def validate_path(path: str) -> tuple[str | None, str | None, str | None]:
    """Resolve a supported Office file inside the configured workspace."""
    if not path:
        return None, None, "path required"
    resolved = os.path.abspath(os.path.expanduser(path))
    workspace_root = workspace_root_from_config()
    if workspace_root is not None:
        if not workspace_root:
            return None, None, "no workspace open"
        root = Path(os.path.abspath(os.path.expanduser(workspace_root))).resolve()
        candidate = Path(resolved).resolve()
        try:
            candidate.relative_to(root)
        except ValueError:
            return None, None, f"path outside workspace: {candidate}"
        resolved = str(candidate)
    if not os.path.isfile(resolved):
        return None, None, f"file not found: {resolved}"
    ext = Path(resolved).suffix.lower()
    kind = EXTENSION_KIND.get(ext)
    if not kind:
        return (
            None,
            None,
            (
                f"unsupported extension {ext!r}; office.* only handles .docx, .xlsx, and .pptx "
                "(use read_file for plain text, image.* for images)"
            ),
        )
    return resolved, kind, None


def resolve_read_source(path: str) -> tuple[str | bytes | None, str | None, str | None]:
    """Resolve a physical workspace file or an exact Runtime attachment route."""
    if not path.startswith("/attachments/"):
        return validate_path(path)
    kind = EXTENSION_KIND.get(Path(path).suffix.lower())
    if kind is None:
        return None, None, f"unsupported attachment extension: {Path(path).suffix.lower()!r}"
    try:
        context = current_runtime_tool_execution().context
    except RuntimeError:
        return None, None, "attachment is not bound to a Runtime tool execution"
    backend = getattr(context, "backend", None)
    if backend is None or not hasattr(backend, "download_files"):
        return None, None, "attachment backend is unavailable"
    response = backend.download_files([path])[0]
    if response.error:
        return None, None, f"attachment is unavailable: {response.error}"
    if response.content is None:
        return None, None, "attachment has no readable content"
    return response.content, kind, None


def resolve_write_target(original_path: str) -> str:
    p = Path(original_path)
    if p.stem.endswith(f".{EDITED_INFIX}"):
        return str(p)
    return str(p.with_name(f"{p.stem}.{EDITED_INFIX}{p.suffix}"))


def ensure_copy_for_write(original_path: str) -> str:
    target = resolve_write_target(original_path)
    if target == original_path:
        return target
    if not os.path.exists(target):
        shutil.copy2(original_path, target)
    return target


def verify_file(path: str, kind: str) -> None:
    if kind == "word":
        DocxDocument(path)
    elif kind == "excel":
        wb = load_workbook(path, read_only=True)
        wb.close()
    elif kind == "powerpoint":
        PptxPresentation(path)


def atomic_write(
    target: str,
    kind: str,
    write_fn: Callable[[str], dict[str, Any]],
) -> dict[str, Any]:
    """Write, validate, and atomically promote a sibling temporary Office file."""
    target_path = Path(target)
    fd, tmp = tempfile.mkstemp(
        suffix=target_path.suffix,
        prefix=f".{target_path.stem}.tmp.",
        dir=str(target_path.parent),
    )
    os.close(fd)
    try:
        summary = write_fn(tmp)
        verify_file(tmp, kind)
        os.replace(tmp, target)
        return summary
    except Exception:
        if os.path.exists(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass
        raise


def write_result(
    original: str,
    target: str,
    kind: str,
    summary: dict[str, Any] | None = None,
) -> dict[str, Any]:
    return {
        "ok": "true",
        "original_path": original,
        "edited_path": target,
        "kind": kind,
        "summary": summary or {},
    }


def write_error(
    kind: str | None,
    message: str,
    original: str | None = None,
) -> dict[str, Any]:
    out: dict[str, Any] = {"ok": "false", "error": message}
    if kind:
        out["kind"] = kind
    if original:
        out["original_path"] = original
    return out
